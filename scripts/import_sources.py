#!/usr/bin/env python3
"""Import real, local TMSA/DRYBMS sources into a private SeedBundle.

Never writes to the source roots, creates accounts, uploads, commits or deletes.
All narrative, source copies and attribution stay beneath private-seed/.
Only a strict-whitelist, empty original-format workbook is made public.
"""
from __future__ import annotations

import argparse
import copy
import io
import hashlib
import json
import re
import shutil
import sys
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.packaging.core import DocumentProperties
from openpyxl.packaging.custom import CustomPropertyList
from openpyxl.worksheet.datavalidation import DataValidationList
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.utils import get_column_letter

# The only authorized scripts are these two files, not a bytecode sidecar.
sys.dont_write_bytecode = True
from verify_sources import bi, expected_response, load, normalize_response, sha, verify

REPO = Path(__file__).resolve().parents[1]
TMSA_ROOT = Path("C:/Users/tuotu/AppData/Local/hermes/cache/scratch/tmsa_consolidation")
DRY_ROOT = Path("C:/Users/tuotu/AppData/Local/hermes/cache/scratch/drybms_tmsa_review")
DEPARTMENTS = {"safety": "安衛處", "technical": "船工處", "crew": "船員組", "management": "管理組", "marine": "海技組", "shipping": "航運處"}
LEVELS = {"Basic": 1, "Intermediate": 2, "Advanced": 3, "Excellence": 4}
EMPTY = {"en": "", "zh": ""}


def write_json(path: Path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def base(entity_id: str, stamp: str):
    # BaseEntity's dates are import dates, not claimed execution/completion dates.
    return {"id": entity_id, "version": 1, "createdAt": stamp, "updatedAt": stamp, "updatedBy": "source-import"}


def literal_locator(record: dict, source_index: dict, field=None, note=""):
    columns = record["columns"]
    if field in {"kpi", "guidance"}:
        col = next(v for k, v in columns.items() if k.startswith("key performance indicator")) if field == "kpi" else columns["best-practice guidance"]
        coord = f"{get_column_letter(col)}{record['row']}"
    elif field == "assessment":
        coord = f"{get_column_letter(columns[record['assessment_header']])}{record['row']}"
    elif field == "evidence":
        col = next((v for k, v in columns.items() if "evidence" in k or ("implement" in k and not k.startswith("2023"))), None)
        coord = f"{get_column_letter(col)}{record['row']}" if col else f"A{record['row']}"
    else:
        coord = f"A{record['row']}:{get_column_letter(max(columns.values()))}{record['row']}"
    return {"source": f"TMSA:{record['source']}", "file": source_index[record["source"]]["filename"],
            "sheet": record["sheet"], "cell": coord, "rawItem": record["raw_item"],
            "note": "; ".join(x for x in [f"Literal source PIC: {record['pic']}", record.get("mapping_warning", ""), note] if x)}


def department_ids(master: dict, primary: list[dict]) -> list[str]:
    # Names embedded in PIC remain literal source data, never roster identities.
    return [key for key, name in DEPARTMENTS.items() if name in master["pic"] or any(p["source"] == key for p in primary)]


def snapshot_sources(private: Path, tmsa_root: Path, dry_root: Path, stamp: str):
    files = []
    for prefix, source_root in [("tmsa", tmsa_root), ("drybms", dry_root)]:
        assert source_root.is_dir(), f"Missing source root: {source_root}"
        for source in sorted(p for p in source_root.rglob("*") if p.is_file()):
            relative = Path("snapshots") / prefix / source.relative_to(source_root)
            target = private / relative
            target.parent.mkdir(parents=True, exist_ok=True)
            fingerprint = sha(source)
            shutil.copy2(source, target)
            assert sha(target) == fingerprint, f"Snapshot copy mismatch: {source}"
            files.append({"sourcePath": str(source.resolve()), "snapshotPath": relative.as_posix(), "bytes": source.stat().st_size, "sha256": fingerprint})
    tmsa_manifest = load(tmsa_root / "manifest.json")
    dry_manifest = load(dry_root / "manifest.json")
    original_manifest = []
    for framework, sources in [("TMSA", tmsa_manifest), ("DRYBMS", [r for r in dry_manifest if r.get("path")])]:
        for record in sources:
            code = record.get("code", record.get("key"))
            original = Path(record["path"])
            assert original.is_file(), f"Original missing (do not normalize its filename): {original}"
            assert sha(original) == record["sha256"], f"Source SHA-256 differs from manifest: {original}"
            assert original.stat().st_size == record["bytes"]
            relative = Path("snapshots/originals") / framework / code / original.name
            target = private / relative
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(original, target)
            assert sha(target) == record["sha256"]
            files.append({"sourcePath": str(original), "snapshotPath": relative.as_posix(), "bytes": record["bytes"], "sha256": record["sha256"]})
            original_manifest.append({**record, "id": f"{framework}:source:{code}", "frameworkId": framework,
                                      "sourceKey": code, "importedAt": stamp, "snapshotPath": relative.as_posix(), "verifiedUnchanged": True})
    assert len(tmsa_manifest) == 10 and len(original_manifest) == 13
    write_json(private / "snapshot-manifest.json", {"schemaVersion": 1, "importedAt": stamp, "files": files})
    return original_manifest


def create_template(master_path: Path, canonical: list[dict], destination: Path, private: Path):
    workbook = load_workbook(master_path, data_only=False, keep_links=False)
    item_positions = {}
    whitelist = {}
    sheets = []
    for ws in workbook:
        records = [r["master"] for r in canonical if r["master"]["sheet"] == ws.title]
        assert records, f"Unexpected master sheet: {ws.title}"
        sample = records[0]
        columns = sample["columns"]
        pic_col = columns["pic"]
        allowed = {c.coordinate: c.value for c in ws[6] if c.column <= pic_col and c.value is not None}
        # E4's column labelled Stage holds the full ID; its Item is the suffix.
        # Preserve both original identifier cells; never replace A7 with stage 1.
        for row in records:
            code = row["item"]
            literal = ws.cell(row["row"], 1).value
            assert literal == code, f"Master ID not in its original A cell: {code}"
            allowed[f"A{row['row']}"] = literal
            if "stage" in columns and "item" in columns:
                suffix = ws.cell(row["row"], columns["item"])
                allowed[suffix.coordinate] = suffix.value
            kpi_col = next(v for k, v in columns.items() if k.startswith("key performance indicator"))
            item_positions[f"TMSA:{code}"] = {"sheet": ws.title, "row": row["row"], "literalItem": code,
                "cells": {"item": f"A{row['row']}", "kpi": f"{get_column_letter(kpi_col)}{row['row']}",
                          "guidance": f"{get_column_letter(columns['best-practice guidance'])}{row['row']}",
                          "response": f"{get_column_letter(columns['y/n'])}{row['row']}",
                          "selfAssessment": f"{get_column_letter(columns[row['assessment_header']])}{row['row']}",
                          "evidence": f"{get_column_letter(next(v for k,v in columns.items() if 'evidence' in k))}{row['row']}",
                          "pic": f"{get_column_letter(pic_col)}{row['row']}"}}
        whitelist[ws.title] = allowed
        sheets.append({"name": ws.title, "chapterId": f"TMSA:{sample['chapter']}", "headerRow": 6,
                       "sourceHeaders": columns, "dataColumns": {"item": 1, "kpi": kpi_col, "guidance": columns["best-practice guidance"],
                       "response": columns["y/n"], "selfAssessment": columns[sample["assessment_header"]],
                       "evidence": next(v for k,v in columns.items() if "evidence" in k), "pic": pic_col},
                       "mergedRanges": list(map(str, ws.merged_cells.ranges)), "printArea": str(ws.print_area),
                       "printTitleRows": ws.print_title_rows, "printTitleCols": ws.print_title_cols,
                       "columnDimensions": {k: {"width": d.width, "min": d.min, "max": d.max, "hidden": d.hidden, "outlineLevel": d.outlineLevel} for k,d in ws.column_dimensions.items()},
                       "rowDimensions": {str(k): {"height": d.height, "hidden": d.hidden, "outlineLevel": d.outlineLevel} for k,d in ws.row_dimensions.items()},
                       "styleArrays": {c.coordinate: list(c._style) if c._style else None for c in ws._cells.values() if not isinstance(c, MergedCell)}})
        # Full _cells inventory includes hidden and out-of-table original content.
        for cell in list(ws._cells.values()):
            if isinstance(cell, MergedCell):
                continue
            if cell.coordinate not in allowed:
                cell.value = None
            cell.comment = None
            cell.hyperlink = None
        for attribute in ["oddHeader", "oddFooter", "evenHeader", "evenFooter", "firstHeader", "firstFooter"]:
            item = getattr(ws, attribute)
            for part in ["left", "center", "right"]:
                getattr(item, part).text = None
        ws.data_validations = DataValidationList()
        ws.conditional_formatting = ConditionalFormattingList()
        ws.tables.clear()
        ws._charts = []
        ws._images = []
        ws.defined_names.clear()
        ws.sheet_properties.codeName = None
    workbook.defined_names.clear()
    workbook._external_links = []
    workbook.properties = DocumentProperties(creator="", lastModifiedBy="", title="", subject="", description="", keywords="", category="", identifier="", language="", version="", contentStatus="")
    workbook.custom_doc_props = CustomPropertyList()
    workbook.code_name = None
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    # openpyxl omits width=0 (falsy) on 630 hidden source dimensions; on
    # reopening that becomes its default 13. Preserve the exact original zero
    # widths in OOXML without reserializing other worksheet XML or namespaces.
    zero_widths = [{(d.min, d.max) for d in ws.column_dimensions.values() if d.width == 0} for ws in workbook]
    archive_bytes = destination.read_bytes()
    with zipfile.ZipFile(io.BytesIO(archive_bytes)) as original_zip:
        with zipfile.ZipFile(destination, "w") as sanitized_zip:
            for info in original_zip.infolist():
                content = original_zip.read(info.filename)
                sheet_match = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", info.filename)
                if sheet_match:
                    widths = zero_widths[int(sheet_match.group(1)) - 1]
                    def preserve_width(match):
                        tag = match.group(0)
                        minimum = re.search(rb'\bmin="(\d+)"', tag)
                        maximum = re.search(rb'\bmax="(\d+)"', tag)
                        if minimum and maximum and (int(minimum.group(1)), int(maximum.group(1))) in widths:
                            if re.search(rb'\bwidth="[^"]*"', tag):
                                return re.sub(rb'\bwidth="[^"]*"', b'width="0"', tag)
                            return tag[:-2] + b' width="0"/>'
                        return tag
                    content = re.sub(rb'<col\b[^>]*/>', preserve_width, content)
                sanitized_zip.writestr(info, content)
    workbook.close()
    assert len(item_positions) == 288 and len(whitelist) == 20
    write_json(private / "tmsa-style-manifest.json", {"schemaVersion": 1, "template": "public/templates/tmsa-template.xlsx",
        "sourceMasterSha256": sha(master_path), "allowedCells": whitelist, "itemPositions": item_positions, "sheets": sheets,
        "reportRules": {"preserveLiteralSheetNames": True, "findRowsByLiteralItemA": True, "bodyFontPt": 15,
        "minimumReportFontPt": 14, "wrapText": True, "longText": "Add labelled continuation rows; never silently cap/crop narratives. Template preserves original geometry; report rendering must be verified separately.",
        "headers": "Only column headers are public; title, chapter-title and AIM narrative bands remain styled and blank. Replace source-year headings explicitly when exporting a current report."}})


def evidence_sources(row: dict, source_area: dict) -> list[dict]:
    result = []
    for cell in source_area["cells"]:
        if cell["page"] not in row["source_pages"]:
            continue
        # Page 35's OCR grouping combined Basic and Intermediate. The original
        # image has an independent Intermediate row, already disclosed by the
        # analytical D17-I-01 record. Split only the literal evidence paragraph;
        # never rename/repair the original source-unit identifiers or body text.
        if source_area["area"] == 17 and cell["page"] == 35 and cell["level"] == "Basic":
            marker = "Verification of:"
            assert cell["suggested_evidence_en"].count(marker) == 1
            start = cell["suggested_evidence_en"].index(marker)
            if row["level"] in {"Basic", "Intermediate"}:
                text = cell["suggested_evidence_en"][:start] if row["level"] == "Basic" else cell["suggested_evidence_en"][start:]
                result.append({**cell, "level": row["level"], "suggested_evidence_en": text,
                    "derivationNote": "Page 35 original-image row boundary: Intermediate starts at 'Verification of:' in the Suggested Objective Evidence column. Original extraction cell is preserved in DRYBMS:original-cells. Exact substring, not invented text."})
        elif cell["level"] == row["level"]:
            result.append(cell)
    return result


def build_bundle(repo: Path, tmsa_root: Path, dry_root: Path, stamp: str):
    private = repo / "private-seed"
    private.mkdir(parents=True, exist_ok=True)
    original_manifest = snapshot_sources(private, tmsa_root, dry_root, stamp)
    tmsa = private / "snapshots/tmsa"
    dry = private / "snapshots/drybms"
    canonical = load(tmsa / "structured/canonical.json")
    registry = load(tmsa / "structured/text_registry.json")
    extra_registry = load(tmsa / "structured/additional_text_registry.json")
    all_registry = {**registry, **extra_registry}
    translations = {}
    for name in [f"translations_{i}.json" for i in range(1,8)] + [f"additional_translations_{i}.json" for i in range(1,4)]:
        for tid, translation in load(tmsa / "out" / name)["translations"].items():
            assert tid not in translations, f"Repeated translation ID: {tid}"
            translations[tid] = translation
    assert set(all_registry) == set(translations)
    source_index = {r["sourceKey"]: r for r in original_manifest if r["frameworkId"] == "TMSA"}
    mapping = load(dry / "analysis/mapping_merged.json")
    units = load(dry / "analysis/drybms_units.json")
    dry_cells = load(dry / "analysis/drybms_cells.json")
    unit_index = {u["unit_id"]: u for u in units}
    dry_source = next(r for r in original_manifest if r["sourceKey"] == "drybms_v12")
    frameworks = [
        {**base("TMSA", stamp), "name": {"en": "TMSA — company checklist", "zh": "TMSA — 公司檢核總表"},
         "description": {"en": "The supplied master is authoritative for KPI/BPG. Source replies are imported, not a newly completed review.", "zh": "以提供的總表作為 KPI／BPG 權威；保留原部門回覆，但匯入不代表已完成本次更新。"},
         "scoring": "tmsa-stage", "active": True, "sourceVersion": "Company master 20250816; supplied department replies 2026",
         "scopeLabel": {"en": "20 company chapters, including Human Factors (14); equally weighted, 0–4. Not an OCIMF official 13-Element score.", "zh": "公司 20 分章（含第 14 章人因），等權平均 0–4；不是 OCIMF 官方 13 Element 分數。"}},
        {**base("DRYBMS", stamp), "name": {"en": "DRYBMS", "zh": "DRYBMS"},
         "description": {"en": "412 analytical rows reference original English source-unit context; that context is not a literal line-by-line translation of the Chinese analysis. Suggested Objective Evidence remains advisory.", "zh": "412 列中文分析引用英文原單元的共享上下文；不是逐列中文的逐字英文翻譯。Suggested Objective Evidence 為建議證據，不升格為標準要求。"},
         "scoring": "drybms-guidance", "active": True, "sourceVersion": "DBCE 2024 V1.2; Guidance v1.1; gap-marked analytical crosswalk",
         "scopeLabel": {"en": "30 equally weighted areas; four percentage self-assessments using Guidance v1.1 p.7. Analytical indices are not official clause IDs.", "zh": "30 主題等權；四級百分比自評按 Guidance v1.1 第 7 頁減權。分析索引不是官方條文編號。"}}
    ]
    provenance = [{"id": "source-import:receipt", "kind": "import-receipt", "importedAt": stamp, "createdBy": "source-import", "updatedBy": "source-import",
        "policies": {"tmsaAuthority": "canonical master KPI/BPG", "latestReplies": "canonical primary (supplementary and template copies preserved separately)",
        "reviewed": False, "executionStatus": "not_started", "namesAreNotAccounts": True, "externalClosure": "unverified", "sourceGapTextGuessed": False,
        "generalThemes": "Explicit improvement areas and element-summary improvement/training themes create unlocated general improvements; strengths, no_finding, confidentiality and response instructions do not."}}]
    for tid, record in all_registry.items():
        provenance.append({"id": f"TMSA:text:{tid}", "kind": "text-registry", "registry": "additional" if tid in extra_registry else "primary",
                           "record": record, "translation": translations[tid], "bilingual": bi(record["text"], tid, all_registry, translations)})
    chapters = []
    items = []
    assessments = []
    seen_chapters = set()
    for sequence, record in enumerate(canonical, 1):
        code = record["item"]
        assert re.fullmatch(r"\d+[A-Z]?\.[1-4]\.\d+", code), f"Invalid authoritative Item literal: {code}"
        master = record["master"]
        text_ids = record["text_ids"]
        cid = f"TMSA:{record['chapter']}"
        if cid not in seen_chapters:
            seen_chapters.add(cid)
            chapters.append({**base(cid, stamp), "frameworkId": "TMSA", "code": record["chapter"],
                "name": bi(master["title"], text_ids["master_title"], all_registry, translations),
                "aim": bi(master["aim"], text_ids["master_aim"], all_registry, translations),
                "order": len(chapters) + 1, "active": True})
        departments = department_ids(master, record["primary"])
        iid = f"TMSA:{code}"
        items.append({**base(iid, stamp), "frameworkId": "TMSA", "chapterId": cid, "code": code,
            "level": int(code.split(".")[1]), "order": sequence, "active": True,
            "kpi": bi(master["kpi"], text_ids["master_kpi"], all_registry, translations),
            "guidance": bi(master["guidance"], text_ids["master_guidance"], all_registry, translations),
            "departmentIds": departments, "sources": [literal_locator(master, source_index, field, "Authoritative master " + field) for field in ["kpi", "guidance"]]})
        selected = record["primary"] if record["primary"] else [master]
        replies = []
        for entry in selected:
            ids = entry.get("text_ids") or {f: text_ids[f"master_{f}"] for f in ["assessment", "evidence"]}
            replies.append({"departmentId": entry["source"], "response": normalize_response(entry["yn"]),
                "assessment": bi(entry["assessment"], ids["assessment"], all_registry, translations),
                "evidence": bi(entry["evidence"], ids["evidence"], all_registry, translations),
                "source": literal_locator(entry, source_index, note=f"Literal Y/N: {entry['yn']}; inherited assessment header: {entry['assessment_header']}; supplied latest reply, not fresh platform review")})
        def joined(field, language):
            if len(replies) == 1:
                return replies[0][field][language]
            return "\n\n".join(f"[{DEPARTMENTS.get(p['departmentId'], p['departmentId']) if language == 'zh' else p['departmentId']} | {p['source']['sheet']}!{p['source']['cell']}]\n{p[field][language]}" for p in replies)
        response, conformity = expected_response(selected)
        note_en = "Imported source assertions only; not freshly reviewed. Assignee, execution status and external finding closure are not inferred from the source prose."
        note_zh = "僅匯入來源自述；尚未完成本次更新。不從正文推定正式負責人帳號、執行狀態或外部缺失已結案。"
        if len({normalize_response(p["yn"]) for p in selected}) > 1:
            note_en += " Distinct departmental Y/N assertions are retained; Yes/No conflict is No/partial, never consensus."
            note_zh += " 已保留部門 Y/N 差異；Yes／No 衝突依 No／部分符合呈現，不製造共識。"
        assessments.append({**base(iid, stamp), "itemId": iid, "response": response, "conformity": conformity,
            "reviewed": False, "selfAssessment": {lang: joined("assessment", lang) for lang in ["en", "zh"]},
            "evidence": {lang: joined("evidence", lang) for lang in ["en", "zh"]}, "departmentIds": departments,
            "assigneeIds": [], "executionStatus": "not_started", "importedYear": 2026, "initialDepartmentResponses": replies,
            "sources": [p["source"] for p in replies], "note": {"en": note_en, "zh": note_zh}})
        provenance.append({"id": f"TMSA:canonical:{code}", "kind": "canonical-source-record", "record": record})
    for area in mapping["areas"]:
        number = area["area"]
        cid = f"DRYBMS:{number}"
        source_area = next(c for c in dry_cells if c["area"] == number)
        chapters.append({**base(cid, stamp), "frameworkId": "DRYBMS", "code": str(number), "name": {"en": source_area["title_en"], "zh": area["title_zh"]}, "order": number, "active": True})
        provenance.append({"id": f"DRYBMS:area:{number}", "kind": "area-analysis", "record": area})
        for row in area["requirements"]:
            iid = f"DRYBMS:{row['rid']}"
            originals = [{"id": uid, "kind": unit_index[uid]["kind"], "en": unit_index[uid]["text_en"], "page": unit_index[uid]["page"]} for uid in row["source_unit_ids"]]
            suggested = evidence_sources(row, source_area)
            assert suggested, f"No literal same-level/page evidence: {row['rid']}"
            source_locators = [{"source": "DRYBMS:drybms_v12", "file": dry_source["path"].split("\\")[-1], "pages": [u["page"]], "rawItem": u["id"],
                "note": f"Original {u['kind']} unit; shared English context, not a literal translation of this Chinese analytical row."} for u in originals]
            source_locators += [{"source": "DRYBMS:drybms_v12", "file": Path(dry_source["path"]).name, "pages": [c["page"]], "rawItem": row["rid"],
                "note": f"Suggested Objective Evidence; {row['level']}; advisory evidence, not a new mandatory standard requirement."} for c in suggested]
            gap_ids = [g["gap_id"] for g in mapping["gaps"] if row["gap_zh"] and g["area"] == number and g["level"] == row["level"] and g["page"] in row["source_pages"]]
            item = {**base(iid, stamp), "frameworkId": "DRYBMS", "chapterId": cid, "code": row["rid"], "level": LEVELS[row["level"]], "order": len([i for i in items if i["frameworkId"] == "DRYBMS"]) + 1,
                "active": True, "kpi": {"en": "\n\n".join(u["en"] for u in originals), "zh": row["requirement_zh"]},
                "guidance": {"en": "\n\n".join(c["suggested_evidence_en"] for c in suggested), "zh": "建議證據原文為英文；參考核對事項：" + row["attention_zh"]},
                "departmentIds": [], "sources": source_locators, "analysisIndex": True, "sourceOriginals": originals,
                "crosswalk": {"tmsaItemCodes": row["tmsa_refs"], "match": row["match"], "difference": row["difference_zh"], "attention": row["attention_zh"]}}
            if row["gap_zh"]:
                item.update(sourceGap=row["gap_zh"], sourceGapResolved=False)
                assert gap_ids, f"Source gap has no linked gap record: {row['rid']}"
            items.append(item)
            assessments.append({**base(iid, stamp), "itemId": iid, "response": None, "conformity": "none", "reviewed": False,
                "selfAssessment": dict(EMPTY), "evidence": dict(EMPTY), "departmentIds": [], "assigneeIds": [], "executionStatus": "not_started", "sources": [],
                "note": {"en": "No current company assessment supplied. Empty/null is not No, Yes or a completed review.", "zh": "未提供本公司的現況自評。空值／null 不是 No、Yes 或已完成更新。"}})
            provenance.append({"id": f"DRYBMS:crosswalk:{row['rid']}", "kind": "analytical-crosswalk", "record": row,
                "sourceOriginals": originals, "suggestedEvidence": suggested, "gapIds": gap_ids})
    historical = load(tmsa / "out/historical_extra_fields.json")
    cycles = load(tmsa / "out/cycle_review.json")
    for pid, kind, data in [
        ("TMSA:registered-sources", "registered-sources", load(tmsa / "registered_sources.json")),
        ("TMSA:source-manifest", "upstream-source-manifest", load(tmsa / "manifest.json")),
        ("TMSA:template-exclusions", "copied-template-references", load(tmsa / "structured/templates_excluded.json")),
        ("TMSA:parsed-source-rows", "source-row-archive", load(tmsa / "structured/source_rows.json")),
        ("TMSA:master-extra-cells", "out-of-table-source-content", load(tmsa / "structured/master_extras.json")),
        ("TMSA:source-coverage", "source-coverage", load(tmsa / "structured/coverage.json")),
        ("TMSA:source-anomalies", "source-anomalies", load(tmsa / "structured/anomalies.json")),
        ("source-import:snapshot-manifest", "durable-snapshot-index", load(private / "snapshot-manifest.json")),
        ("TMSA:historical-extra-fields", "historical-fields", historical),
        ("TMSA:annotations", "source-annotations", load(tmsa / "out/annotation_translations.json")),
        ("TMSA:cycle-review-metadata", "cycle-metadata", {"notes": cycles["notes"], "coverage": cycles["coverage"]}),
        ("DRYBMS:registered-sources", "upstream-source-manifest", load(dry / "manifest.json")),
        ("DRYBMS:original-units", "original-units", units), ("DRYBMS:original-cells", "original-cells", dry_cells),
        ("DRYBMS:gaps", "source-gaps", mapping["gaps"]), ("DRYBMS:excluded-units", "excluded-units", mapping["excluded_units"]),
        ("DRYBMS:coverage", "upstream-verification", load(dry / "analysis/coverage_final.json")),
        ("DRYBMS:verification", "upstream-verification", load(dry / "verification.json"))]:
        provenance.append({"id": pid, "kind": kind, "record": data})
    for historical_field in historical:
        assert historical_field["text_id"] in all_registry
    inspections = []
    improvements = []
    finding_links = []
    tmsa_items = [i for i in items if i["frameworkId"] == "TMSA"]
    item_index = {i["id"]: i for i in items}
    report_sets = [("mandeep", ["audit_mandeep_early.json", "audit_mandeep_late.json"], "2025-05-23", "Mandeep / Quantum Vetting — 19–23 May 2025", "Quantum Vetting"),
                   ("exxon", ["audit_exxon.json"], "2025-09-09", "ExxonMobil — 09 Sep 2025", "ExxonMobil / IMT")]
    for code, report_files, date, title, organization in report_sets:
        inspection_id = f"inspection:{code}:2025"
        findings = []
        dedupes = set()
        for report_file in report_files:
            report = load(tmsa / "out" / report_file)
            provenance.append({"id": f"TMSA:report:{report_file}", "kind": "inspection-source-report", "record": report})
            for record_kind, entries in [("section", report["sections"]), ("general", report["general"])]:
                for entry in entries:
                    # Literal source content + exact reported identifier/location bind replay.
                    source_identity = {"source": code, "kind": record_kind, "record": entry}
                    suffix = hashlib.sha256(json.dumps(source_identity, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")).hexdigest()[:24]
                    dedupe_key = f"TMSA:{code}:{record_kind}:{suffix}"
                    assert dedupe_key not in dedupes, f"Repeated external source record: {dedupe_key}"
                    dedupes.add(dedupe_key)
                    finding_id = f"{inspection_id}:finding:{suffix}"
                    if record_kind == "section":
                        iid = f"TMSA:{entry['item']}"
                        assert iid in item_index, f"Unresolved reported Item must not be guessed: {entry['item']}"
                        item_ids = [iid]
                        classification = entry["status"]
                        observation = {"en": entry["observation_en"], "zh": entry["observation_zh"]}
                        recommendation = {"en": entry["recommendation_en"], "zh": entry["recommendation_zh"]}
                        heading = entry["reported_heading"]
                        should_improve = classification == "finding"
                        finding_title = f"{code} | {heading}"
                    else:
                        item_ids = []
                        classification = "positive" if entry["category"] == "strength" else "general" if entry["category"] in {"improvement_area", "element_summary"} else "guidance"
                        observation = {"en": entry["text_en"], "zh": entry["text_zh"]}
                        recommendation = dict(EMPTY)
                        heading = entry["title_en"]
                        should_improve = entry["category"] in {"improvement_area", "element_summary"}
                        finding_title = f"{code} | {entry['title_en']} / {entry['title_zh']}"
                    warning = "\n\n".join(x for x in [entry.get("mapping_warning", ""), entry.get("mapping_warning_zh", "")] if x)
                    source_locator = {"source": f"TMSA:{code}", "file": source_index[code]["filename"], "pages": entry["pdf_pages"], "rawItem": heading,
                        "note": f"Source record: {report_file}; {record_kind}; classification/category: {entry.get('classification', entry.get('category', entry.get('status', '')))}; printed pages: {entry.get('printed_pages', [])}."}
                    finding = {"id": finding_id, "itemIds": item_ids, "classification": classification, "observation": observation, "recommendation": recommendation, "sources": [source_locator]}
                    if warning:
                        finding["warning"] = warning
                    findings.append(finding)
                    finding_links.append({"sourceKey": code, "recordKind": record_kind, "record": entry, "inspectionId": inspection_id, "findingId": finding_id,
                                          "dedupeKey": dedupe_key, "createsImprovement": should_improve})
                    if should_improve:
                        narrative = {}
                        for language in ["en", "zh"]:
                            obs_title = "Observation" if language == "en" else "觀察"
                            rec_title = "Recommendation" if language == "en" else "建議"
                            narrative[language] = f"{obs_title}:\n{observation[language]}" + (f"\n\n{rec_title}:\n{recommendation[language]}" if recommendation[language] else "")
                        improvements.append({**base(f"improvement:{suffix}", stamp), "title": finding_title, "content": narrative, "itemIds": item_ids,
                            "departmentIds": sorted({d for iid in item_ids for d in item_index[iid]["departmentIds"]}), "assigneeIds": [], "source": "external",
                            "inspectionId": inspection_id, "findingId": finding_id, "sourceKey": dedupe_key, "urgency": "medium", "status": "open", "percent": 0,
                            "targetDate": "", "plan": "", "events": [], "closedAt": None, "closedBy": None,
                            "sources": [{**source_locator, "note": source_locator["note"] + " Closure unverified / 結案未核實; priority is an editable import default, not an auditor-assigned rating."}]})
        result_items = []
        for item in tmsa_items:
            related = [f for f in findings if item["id"] in f["itemIds"]]
            result_items.append({"itemId": item["id"], "response": None, "conformity": "none", "reviewed": False,
                "observation": {lang: "\n\n".join(f["observation"][lang] for f in related if f["observation"][lang]) for lang in ["en", "zh"]},
                "recommendation": {lang: "\n\n".join(f["recommendation"][lang] for f in related if f["recommendation"][lang]) for lang in ["en", "zh"]}})
        inspections.append({**base(inspection_id, stamp), "frameworkId": "TMSA", "title": title, "organization": organization, "date": date, "type": "formal", "status": "final",
            "snapshot": {"framework": copy.deepcopy(frameworks[0]), "chapters": copy.deepcopy([c for c in chapters if c["frameworkId"] == "TMSA"]), "items": copy.deepcopy(tmsa_items), "scoringVersion": "company-20-chapters-contiguous-stage-v1"},
            "results": result_items, "findings": findings, "imported": True,
            "scoreNote": {"en": "Only the findings/feedback report is available, not a complete graded per-item score sheet. All 288 snapshot results remain unassessed (null/unreviewed); unreported or no-finding Items are not Yes and OFIs do not automatically mean No. Snapshot scope includes all 20 company chapters, not a claim that the auditor assessed every chapter. Date is the source report/assessment end date; import timestamp is separate.",
                          "zh": "只有發現／回饋報告，沒有完整逐條成績。288 項快照結果全部未評（null／未更新）；未報項目或無發現不等於 Yes，OFI 不自動判 No。快照保存公司完整 20 分章，不聲稱稽核已逐章完成評定。日期為來源報告／稽核結束日；匯入時間另記。"}})
    provenance.append({"id": "TMSA:external-finding-links", "kind": "external-finding-ledger", "record": finding_links})
    dry_scores = [{**base(f"DRYBMS:{a['area']}", stamp), "chapterId": f"DRYBMS:{a['area']}", "levels": [None, None, None, None],
                   "note": "未提供四級百分比自評；不從中文對照或外部發現推定評分。 / No supplied four-level percentage assessment; no scores inferred from crosswalks or findings."} for a in mapping["areas"]]
    bundle = {"schemaVersion": 1, "frameworks": frameworks, "chapters": chapters, "items": items, "assessments": assessments,
        "improvements": improvements, "inspections": inspections, "dryScores": dry_scores, "cycles": cycles["rows"], "provenance": provenance,
        "sourceManifest": original_manifest, "importNote": {"en": f"Private local import at {stamp}. Source paragraphs, conflicting responses, historical extra fields, original identifiers, annotations, cycle quotes and analysis gaps retained. No credentials/accounts or completion/closure/score claims created.",
        "zh": f"私有本機匯入時間：{stamp}。保留完整來源段落、回覆衝突、歷史附加欄、原編號、批註、週期引用及分析缺口。不建立憑證／人員帳號，也不捏造更新完成、結案或評分。"}}
    write_json(private / "platform-seed.json", bundle)
    master_path = private / source_index["master"]["snapshotPath"]
    create_template(master_path, canonical, repo / "public/templates/tmsa-template.xlsx", private)
    return bundle


def main():
    if not __debug__:
        raise RuntimeError("Source import must run without Python -O/-OO (verification assertions enabled).")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--repo", type=Path, default=REPO)
    parser.add_argument("--tmsa-root", type=Path, default=TMSA_ROOT)
    parser.add_argument("--dry-root", type=Path, default=DRY_ROOT)
    parser.add_argument("--native-excel", action="store_true")
    args = parser.parse_args()
    stamp = datetime.now(timezone.utc).isoformat()
    try:
        bundle = build_bundle(args.repo.resolve(), args.tmsa_root.resolve(), args.dry_root.resolve(), stamp)
        report = verify(args.repo.resolve(), False, args.native_excel)
        write_json(args.repo / "private-seed/source-verification.json", report)
        print(json.dumps({"passed": report["passed"], "importedAt": stamp, "counts": report["counts"],
                          "seed": str((args.repo / "private-seed/platform-seed.json").resolve()), "template": str((args.repo / "public/templates/tmsa-template.xlsx").resolve()),
                          "verification": str((args.repo / "private-seed/source-verification.json").resolve())}, ensure_ascii=False, indent=2))
        return 0
    except (AssertionError, OSError, KeyError, ValueError, StopIteration) as error:
        import traceback
        frame = traceback.extract_tb(error.__traceback__)[-1]
        detail = f"{type(error).__name__}: {str(error) or frame.line} ({Path(frame.filename).name}:{frame.lineno})"
        report = {"passed": False, "importedAt": stamp, "errors": [detail]}
        write_json(args.repo / "private-seed/source-verification.json", report)
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 1


if __name__ == "__main__":
    sys.exit(main())
