"""Verify saved original-browser remarks downloads, including read-only native Excel.
Run: uv run --offline --no-project --python 3.12 --with pymupdf==1.26.5 --with openpyxl==3.1.5 --with pywin32 python scripts/verify-item-remarks-artifacts.py RECEIPT
Synthetic evidence only. No hosted access or source-workbook changes.
"""
from pathlib import Path
from collections import Counter
import gc
import hashlib
import json
import re
import shutil
import sys
import subprocess
import unicodedata
import uuid
import fitz
import openpyxl

STATUS = [
    ('fully_met', '完全符合', 'DFF2E6', '195C36'),
    ('partial_improvable', '部分符合可完善', 'E1EFFA', '1F5578'),
    ('partial_blocked', '部分符合近期無法完善', 'EEE5F7', '643881'),
    ('not_met_improvable', '不符合可完善', 'FCE8CD', '8A4909'),
    ('not_met_blocked', '不符合近期無法完善', 'FBE1E1', '9B2929'),
]

def compact(text):
    return re.sub(r'\s+', '', unicodedata.normalize('NFKC', text))

def spans(page):
    return [s for b in page.get_text('dict')['blocks'] for line in b.get('lines', []) for s in line['spans']]

def check_text(document, expected, content_x, header):
    content = ''
    all_text = '\n'.join(p.get_text() for p in document)
    for page in document:
        ordered = sorted(spans(page), key=lambda s: (round(s['origin'][1], 2), s['origin'][0]))
        content += ''.join(s['text'] for s in ordered if s['bbox'][0] >= content_x-1 and compact(s['text']) != compact(header))
    actual = compact(content)
    for text, count in Counter(compact(n['remarks']) for n in expected if n['remarks']).items():
        assert actual.count(text) == count, f'Full remarks missing, duplicated or clipped: {text[:40]}'
    for n in expected:
        assert n['code'] in all_text, f'Missing literal Item: {n["code"]}'
    assert 'DO_NOT_EXPORT_SOURCE_NOTE' not in all_text
    assert 'OUTSIDE_SELECTED_CHAPTER' not in all_text
    assert '未選章節' not in all_text
    assert 'TMSA-R7' not in all_text and 'DRYBMS-R7' not in all_text
    assert '1.1.3' not in all_text, 'Status-only item must not be exported'
    return all_text

def verify_pdf(path, entry):
    doc = fitz.open(path)
    x = 34+(1190.55-68)*.53+8
    check_text(doc, entry['expected'], x, '備註內容')
    fills, text_colours, sizes = [], set(), set()
    for page in doc:
        for s in spans(page):
            if not s['text'].strip():
                continue
            sizes.add(round(s['size'], 2))
            text_colours.add(s['color'])
            assert s['size'] >= 13.99, ('Font below 14pt', s)
            box = s['bbox']
            assert box[0] >= 33 and box[2] <= page.rect.width-33, ('Horizontal overflow', s)
            assert box[1] >= 0 and box[3] <= page.rect.height, ('Page overflow', s)
        fills.extend(d['fill'] for d in page.get_drawings() if d['fill'] is not None)
    for _, label, bg, fg in STATUS:
        assert compact(label) in compact('\n'.join(p.get_text() for p in doc))
        colour = tuple(int(bg[n:n+2], 16)/255 for n in (0, 2, 4))
        assert any(all(abs(a-b) < .005 for a, b in zip(colour, f)) for f in fills), f'Missing status fill: {label}'
        assert int(fg, 16) in text_colours, f'Missing status text colour: {label}'
    assert len(doc) > 1, 'Long remarks did not exercise continuation pages'
    for index in {0, 1, len(doc)-1}:
        doc[index].get_pixmap(matrix=fitz.Matrix(1.2, 1.2), alpha=False).save(str(path.with_name(path.stem+f'-page-{index+1}.png')))
    return {'pages': len(doc), 'fontSizes': sorted(sizes), 'fullTextAndScope': 'PASS', 'fiveStatusColours': 'PASS'}

def verify_excel(path, entry):
    book = openpyxl.load_workbook(path, data_only=False)
    assert len(book.worksheets) == 1
    sheet = book.active
    rows = list(sheet.iter_rows(min_row=5, max_col=4))
    assert {r[1].value for r in rows} == {n['code'] for n in entry['expected']}
    for item in entry['expected']:
        own = [r for r in rows if r[1].value == item['code']]
        assert ''.join(str(r[3].value or '') for r in own) == item['remarks']
        status = next(s for s in STATUS if s[0] == item['remarksStatus'])
        for row in own:
            assert row[2].value == status[1]
            assert row[2].fill.fgColor.rgb == 'FF'+status[2]
            assert row[2].font.color.rgb == 'FF'+status[3]
    for row in range(4, sheet.max_row+1):
        assert sheet.row_dimensions[row].height <= 409.5
        for col in range(1, 5):
            c = sheet.cell(row, col)
            assert c.data_type != 'f', 'User text must never become a formula'
            assert c.alignment.wrap_text
            assert all(getattr(c.border, name).style for name in ('left', 'right', 'top', 'bottom'))
    return {'uniqueItems': len(entry['expected']), 'physicalRows': len(rows), 'completeText': 'PASS', 'coloursAndBorders': 'PASS'}

def native_excel_worker(path, entry):
    import win32com.client
    import win32process
    import win32api
    import win32con
    import win32event
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    owned = path.with_name(path.stem+'-owned-'+uuid.uuid4().hex[:8]+'.xlsx')
    printed = owned.with_suffix('.pdf')
    shutil.copy2(path, owned)
    existing_pids=set(win32process.EnumProcesses())
    app = win32com.client.DispatchEx('Excel.Application')
    app.Visible = False
    app.DisplayAlerts = False
    app.EnableEvents = False
    app.AutomationSecurity = 3
    _, pid = win32process.GetWindowThreadProcessId(app.Hwnd)
    handle = win32api.OpenProcess(win32con.SYNCHRONIZE | win32con.PROCESS_TERMINATE, False, pid)
    book = sheet = cell = None
    result = {'ownedPid': pid, 'newOwnedProcess': pid not in existing_pids, 'forcedOwnedStop': False}
    try:
        book = app.Workbooks.Open(str(owned.resolve()), UpdateLinks=0, ReadOnly=True, CorruptLoad=0)
        assert book.ReadOnly and book.Worksheets.Count == 1
        sheet = book.Worksheets(1)
        assert sheet.PageSetup.Zoom is False and sheet.PageSetup.FitToPagesWide == 1
        native_rows = []
        for row in range(5, sheet.UsedRange.Rows.Count+1):
            values = [sheet.Cells(row, c).Value for c in range(1, 5)]
            native_rows.append(values)
            for col in range(1, 5):
                cell = sheet.Cells(row, col)
                assert cell.WrapText and not cell.HasFormula
                assert all(cell.Borders(edge).LineStyle != -4142 for edge in (7, 8, 9, 10))
        assert {r[1] for r in native_rows} == {n['code'] for n in entry['expected']}
        for n in entry['expected']:
            assert ''.join(r[3] or '' for r in native_rows if r[1] == n['code']) == n['remarks']
        book.ExportAsFixedFormat(0, str(printed.resolve()))
        assert printed.is_file()
        result.update(acceptedWithCorruptLoad0=True, readOnly=True, nativePdf=str(printed), copy=str(owned))
    finally:
        cell = sheet = None
        if book is not None:
            book.Close(SaveChanges=False)
        book = None
        app.Quit()
        app = None
        gc.collect()
        result['ownedExcelStopped'] = win32event.WaitForSingleObject(handle, 8000) == win32event.WAIT_OBJECT_0
        if not result['ownedExcelStopped']:
            assert result['newOwnedProcess'], 'Will not stop a pre-existing Excel process'
            # The synthetic workbook was closed without save; stop ONLY this fresh owned instance.
            win32api.TerminateProcess(handle, 0)
            result['forcedOwnedStop']=True
            result['ownedExcelStopped']=win32event.WaitForSingleObject(handle, 5000)==win32event.WAIT_OBJECT_0
        handle.Close()
        # Excel can keep COM proxy references until this owned Python worker exits.
        result['stoppedBeforeWorkerExit'] = result.pop('ownedExcelStopped')
        result['quitRequested'] = True
        owned.with_suffix('.native.json').write_text(json.dumps(result,ensure_ascii=True,indent=2),encoding='utf-8')
        assert hashlib.sha256(path.read_bytes()).hexdigest() == before
    doc = fitz.open(printed)
    headers = doc[0].search_for('備註內容')
    assert headers, 'Native Excel print is missing the remarks column'
    check_text(doc, entry['expected'], headers[0].x0, '備註內容')
    for index in {0, len(doc)-1}:
        doc[index].get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).save(str(printed.with_name(printed.stem+f'-page-{index+1}.png')))
    result.update(nativePrintPages=len(doc), nativePrintFullText='PASS')
    return result

def native_excel(path, entry):
    import win32api
    import win32con
    import win32event
    # Close the COM-owning process before checking Excel process lifetime.
    worker = subprocess.run([sys.executable, str(Path(__file__).resolve()), '--native'], input=json.dumps(entry), text=True, encoding='utf-8', capture_output=True, timeout=120)
    assert worker.returncode == 0, worker.stderr
    result = json.loads(worker.stdout)
    try:
        handle = win32api.OpenProcess(win32con.SYNCHRONIZE, False, result['ownedPid'])
    except Exception as exc:
        assert getattr(exc, 'winerror', None) == 87, str(exc)
        result['ownedExcelStopped'] = True
    else:
        result['ownedExcelStopped'] = win32event.WaitForSingleObject(handle, 8000) == win32event.WAIT_OBJECT_0
        handle.Close()
    assert result['ownedExcelStopped'], 'Owned Excel remained after its COM worker exited'
    return result

def main():
    receipt_path = Path(sys.argv[1])
    receipt = json.loads(receipt_path.read_text(encoding='utf-8'))
    result = {'result': 'RUNNING', 'sourceReceipt': str(receipt_path.resolve()), 'artifacts': [], 'verifierSha256': hashlib.sha256(Path(__file__).read_bytes()).hexdigest()}
    target = receipt_path.parent / 'artifact-verification.json'
    try:
        entries = receipt['downloads']
        assert len(entries) == 4
        assert {(e['framework'], e['format']) for e in entries} == {(f, t) for f in ('TMSA', 'DRYBMS') for t in ('xlsx', 'pdf')}
        for entry in sorted(entries, key=lambda e: (e['format'] != 'pdf', e['framework'])):
            path = Path(entry['path'])
            assert hashlib.sha256(path.read_bytes()).hexdigest() == entry['sha256']
            checks = verify_excel(path, entry) if entry['format'] == 'xlsx' else verify_pdf(path, entry)
            if entry['format'] == 'xlsx':
                checks['nativeExcel'] = native_excel(path, entry)
            result['artifacts'].append({'path': str(path), 'sha256': entry['sha256'], **checks})
            target.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
        result['result'] = 'PASS'
    except Exception as exc:
        result['result'] = 'FAIL'
        result['failure'] = str(exc)
        raise
    finally:
        target.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
        print(json.dumps({'result': result['result'], 'verifiedArtifacts': len(result['artifacts']), 'receipt': str(target)}, ensure_ascii=True))

if __name__ == '__main__':
    if sys.argv[1:]==['--native']:
        entry=json.loads(sys.stdin.read())
        print(json.dumps(native_excel_worker(Path(entry['path']),entry),ensure_ascii=True))
    else:
        main()
