#!/usr/bin/env python3
"""Local-only source verification. Real inputs are required; no fixture seed.

Run with the pinned environment described in docs/SOURCE_DATA.md.
The default verification binds the durable snapshots AND live originals.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
import subprocess
import sys
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
import pymupdf

DEFAULT_REPO = Path(__file__).resolve().parents[1]


def load(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def sha(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def bi(raw: str, text_id, registry: dict, translations: dict) -> dict:
    if not raw.strip():
        return {"en": "", "zh": ""}
    assert text_id in registry, f"Unregistered paragraph: {text_id}"
    assert registry[text_id]["text"] == raw, f"Literal registry mismatch: {text_id}"
    t = translations[text_id]
    assert t.get("zh"), f"Missing translation: {text_id}"
    if "en" not in t:
        assert not re.search(r"[\u3400-\u9fff]", raw), f"Mixed source lacks supplied English: {text_id}"
    return {"en": t.get("en", raw), "zh": t["zh"]}


def normalize_response(raw: str):
    # Keep the literal token separately; this only maps the response enum.
    token = raw.strip().upper()
    return "yes" if token in {"Y", "YES"} else "no" if token in {"N", "NO"} else None


def expected_response(records: list[dict]):
    responses = [normalize_response(r["yn"]) for r in records]
    distinct = set(responses)
    if "yes" in distinct and "no" in distinct:
        return "no", "partial"
    if "no" in distinct:
        return "no", "none"
    if distinct == {"yes"}:
        return "yes", "all"
    return None, "none"


def compact(text: str) -> str:
    return re.sub(r"\s+", "", text)


def base_check(entity: dict, imported_at: str):
    assert entity["version"] == 1
    assert entity["updatedBy"] == "source-import"
    assert entity["createdAt"] == entity["updatedAt"] == imported_at
    if "createdBy" in entity:
        assert entity["createdBy"] == "source-import"


def verify_pdf_columns(source_path: Path, areas: list[dict], units: list[dict]) -> int:
    """Re-extract every saved source column using its literal page/y locator.

    Only whitespace/line grouping from the established extraction is applied.
    OCR spellings, fragments, exclusions and the page-35 grouping are retained.
    """
    doc = pymupdf.open(source_path)
    checked = 0
    reconstructed_units = []
    for area in areas:
        for cell in area["cells"]:
            words = doc[cell["page"] - 1].get_text("words")
            header = next(w for w in words if w[0] < 85 and 100 < w[1] < 160 and re.sub("[^a-z]", "", w[4].lower()) in {"level", "leve", "leel", "levei"})
            x = header[0]
            bounds = [x - 8.5, x + 70.5, x + 190.9, x + 333.5, x + 489.5]
            y0, y1 = cell["y_range"]
            for index, field in enumerate(["expectations_en", "targets_en", "suggested_evidence_en"], 1):
                chosen = [w for w in words if bounds[index] < (w[0] + w[2]) / 2 < bounds[index + 1] and y0 <= (w[1] + w[3]) / 2 < y1]
                lines = []
                for word in sorted(chosen, key=lambda w: (w[1], w[0])):
                    if not lines or abs(lines[-1][0] - word[1]) > 4:
                        lines.append([word[1], [word]])
                    else:
                        lines[-1][1].append(word)
                fragments = []
                previous_y = None
                for y, line in lines:
                    if previous_y is not None and y - previous_y > 18:
                        fragments.append("")
                    fragments.append(" ".join(w[4] for w in sorted(line, key=lambda w: w[0])))
                    previous_y = y
                literal = "\n".join(fragments).strip()
                assert literal == cell[field], f"Nonliteral PDF column: area {area['area']} page {cell['page']} {cell['level']} {field}"
                checked += 1
                if field != "suggested_evidence_en":
                    kind = "E" if field == "expectations_en" else "T"
                    level = {"Basic": "B", "Intermediate": "I", "Advanced": "A", "Excellence": "X"}[cell["level"]]
                    for number, paragraph in enumerate([s.strip() for s in re.split(r"\n\s*\n", literal) if s.strip()], 1):
                        reconstructed_units.append({"unit_id": f"D{area['area']:02d}-{level}-p{cell['page']}-{kind}{number}",
                            "area": area["area"], "level": cell["level"], "kind": "Expectation" if kind == "E" else "Target",
                            "page": cell["page"], "text_en": paragraph, "source_y_range": cell["y_range"]})
    doc.close()
    assert reconstructed_units == units, "Original 570 source-unit paragraphs differ from the hash-bound PDF"
    return checked


def native_excel(template: Path) -> dict:
    # Owned Excel instance only; no user workbook is saved, closed or modified.
    escaped = template.resolve().as_posix().replace("'", "''")
    script = """
[Console]::OutputEncoding = [Text.UTF8Encoding]::new($false)
$ErrorActionPreference = 'Stop'
$ProgressPreference = 'SilentlyContinue'
$excel = $null; $book = $null
try {
  $excel = New-Object -ComObject Excel.Application
  $excel.Visible = $false
  $excel.DisplayAlerts = $false
  $excel.AskToUpdateLinks = $false
  $excel.AutomationSecurity = 3
  $nativePath = [IO.Path]::GetFullPath('%s')
  # Open's normal CorruptLoad=0 path, never a repair/recovery mode.
  $book = $excel.Workbooks.Open($nativePath,0,$true)
  $sheets = @()
  foreach ($sheet in $book.Worksheets) {
    $sheets += [pscustomobject]@{name=$sheet.Name; visible=$sheet.Visible; usedRange=$sheet.UsedRange.Address(); item=$sheet.Cells.Item(7,1).Value2; header=$sheet.Cells.Item(6,1).Value2; body=$sheet.Cells.Item(7,3).Value2}
  }
  [pscustomobject]@{verified=$true; application='Microsoft Excel'; version=$excel.Version; readOnly=$book.ReadOnly; worksheetCount=$book.Worksheets.Count; sheets=$sheets; sourceSaved=$false; corruptLoad=0} | ConvertTo-Json -Depth 8 -Compress
} finally {
  if ($null -ne $book) {$book.Close($false); [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($book)}
  if ($null -ne $excel) {$excel.Quit(); [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($excel)}
}
""" % escaped
    import base64
    result = subprocess.run(
        ["powershell.exe", "-NoProfile", "-NonInteractive", "-EncodedCommand",
         base64.b64encode(script.encode("utf-16le")).decode("ascii")],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120,
    )
    assert result.returncode == 0, f"Native Excel failed: {result.stderr}"
    report = json.loads(result.stdout.strip().lstrip("\ufeff"))
    assert report["readOnly"] and report["worksheetCount"] == 20
    assert all(x["body"] is None for x in report["sheets"])
    return report


def verify(repo: Path, snapshots_only=False, run_native=False, seed_path: Path | None = None, template_path: Path | None = None) -> dict:
    if not __debug__:
        raise RuntimeError("Verification must run without Python -O/-OO (assertions enabled).")
    private = repo / "private-seed"
    path = seed_path or private / "platform-seed.json"
    assert path.is_file(), f"Missing real import artifact: {path}"
    seed = load(path)
    assert seed["schemaVersion"] == 1
    receipt = next(p for p in seed["provenance"] if p["kind"] == "import-receipt")
    stamp = receipt["importedAt"]
    assert datetime.fromisoformat(stamp.replace("Z", "+00:00")).tzinfo is not None
    prov = {p["id"]: p for p in seed["provenance"]}
    tmsa = private / "snapshots" / "tmsa"
    dry = private / "snapshots" / "drybms"
    canonical = load(tmsa / "structured/canonical.json")
    mapping = load(dry / "analysis/mapping_merged.json")
    units = load(dry / "analysis/drybms_units.json")
    cells = load(dry / "analysis/drybms_cells.json")
    registry = load(tmsa / "structured/text_registry.json")
    extra_registry = load(tmsa / "structured/additional_text_registry.json")
    all_registry = {**registry, **extra_registry}
    translations = {}
    for name in [f"translations_{i}.json" for i in range(1, 8)] + [f"additional_translations_{i}.json" for i in range(1, 4)]:
        for key, value in load(tmsa / "out" / name)["translations"].items():
            assert key not in translations, f"Duplicate translated ID: {key}"
            translations[key] = value
    assert set(all_registry) == set(translations)
    assert len(registry) == 1451 and len(extra_registry) == 111
    for tid, r in all_registry.items():
        actual = prov[f"TMSA:text:{tid}"]
        assert actual["record"] == r, f"Registry/use omission: {tid}"
        assert actual["translation"] == translations[tid]
        assert actual["bilingual"] == bi(r["text"], tid, all_registry, translations)
    assert {p["id"] for p in seed["provenance"] if p["kind"] == "text-registry"} == {f"TMSA:text:{x}" for x in all_registry}

    for key, expected in [("frameworks", 2), ("chapters", 50), ("items", 700), ("assessments", 700), ("inspections", 2), ("dryScores", 30)]:
        records = seed[key]
        assert len(records) == len({r["id"] for r in records}) == expected, f"Wrong unique {key}: {len(records)}"
        for entity in records:
            base_check(entity, stamp)
    for entity in seed["improvements"]:
        base_check(entity, stamp)
    frameworks = {f["id"]: f for f in seed["frameworks"]}
    assert set(frameworks) == {"TMSA", "DRYBMS"}
    assert frameworks["TMSA"]["scoring"] == "tmsa-stage"
    assert frameworks["DRYBMS"]["scoring"] == "drybms-guidance"
    chapters = {c["id"]: c for c in seed["chapters"]}
    items = {i["id"]: i for i in seed["items"]}
    assessments = {a["id"]: a for a in seed["assessments"]}
    assert set(items) == set(assessments)
    assert all(i["chapterId"] in chapters and i["active"] for i in items.values())
    assert all(a["itemId"] == a["id"] and a["assigneeIds"] == [] for a in assessments.values())
    assert all(not a["reviewed"] and a["executionStatus"] == "not_started" for a in assessments.values()), "Import is not a fresh company review"
    known_departments = {"safety", "technical", "crew", "management", "marine", "shipping"}
    assert all(set(i["departmentIds"]) <= known_departments for i in items.values())
    assert all(set(a["departmentIds"]) <= known_departments for a in assessments.values())

    live_workbooks = {}
    original_hashes = []
    upstream_sources = {("TMSA", r["code"]): r for r in load(tmsa / "manifest.json")}
    upstream_sources.update({("DRYBMS", r["key"]): r for r in load(dry / "manifest.json") if r.get("path")})
    assert len(seed["sourceManifest"]) == 13
    for source in seed["sourceManifest"]:
        original_record = upstream_sources[(source["frameworkId"], source["sourceKey"])]
        assert all(source[k] == value for k, value in original_record.items()), "Original manifest/path/filename tokens changed"
        copied = private / source["snapshotPath"]
        assert sha(copied) == source["sha256"]
        assert copied.stat().st_size == source["bytes"]
        if not snapshots_only:
            original = Path(source["path"])
            assert original.is_file(), f"Missing original: {original}"
            assert sha(original) == source["sha256"], f"Changed original: {source['id']}"
        original_hashes.append({"id": source["id"], "sha256": source["sha256"], "liveUnchanged": not snapshots_only, "snapshotIdentical": True})
        if source["frameworkId"] == "TMSA" and copied.suffix.lower() == ".xlsx":
            live_workbooks[source["sourceKey"]] = load_workbook(copied, data_only=False, keep_links=False)
    assert Counter(s["frameworkId"] for s in seed["sourceManifest"]) == {"TMSA": 10, "DRYBMS": 3}
    literal_cells = 0
    conflicts = []
    chapter_counts = Counter()
    for record in canonical:
        code = record["item"]
        assert re.fullmatch(r"\d+[A-Z]?\.[1-4]\.\d+", code), f"Invalid literal master token: {code}"
        iid = f"TMSA:{code}"
        actual = items[iid]
        chapter_counts[record["chapter"]] += 1
        assert actual["code"] == code and actual["chapterId"] == f"TMSA:{record['chapter']}"
        assert actual["level"] == int(code.split(".")[1])
        assert prov[f"TMSA:canonical:{code}"]["record"] == record
        master = record["master"]
        for field, target in [("kpi", "kpi"), ("guidance", "guidance")]:
            assert actual[target] == bi(master[field], record["text_ids"][f"master_{field}"], all_registry, translations)
        selected = record["primary"] if record["primary"] else [master]
        expected_yn, expected_conformity = expected_response(selected)
        assessment = assessments[iid]
        assert (assessment["response"], assessment["conformity"]) == (expected_yn, expected_conformity)
        if {normalize_response(p["yn"]) for p in selected} >= {"yes", "no"}:
            conflicts.append(code)
        assert len(assessment["initialDepartmentResponses"]) == len(selected)
        for entry, response in zip(selected, assessment["initialDepartmentResponses"]):
            assert response["response"] == normalize_response(entry["yn"])
            text_ids = entry.get("text_ids") or {f: record["text_ids"][f"master_{f}"] for f in ["assessment", "evidence"]}
            assert response["assessment"] == bi(entry["assessment"], text_ids["assessment"], all_registry, translations)
            assert response["evidence"] == bi(entry["evidence"], text_ids["evidence"], all_registry, translations)
            assert response["source"]["rawItem"] == entry["raw_item"]
            assert response["departmentId"] == entry["source"]
        department_names = {"safety": "安衛處", "technical": "船工處", "crew": "船員組", "management": "管理組", "marine": "海技組", "shipping": "航運處"}
        for field, target in [("assessment", "selfAssessment"), ("evidence", "evidence")]:
            for language in ["en", "zh"]:
                replies = assessment["initialDepartmentResponses"]
                full_text = replies[0][field][language] if len(replies) == 1 else "\n\n".join(f"[{department_names.get(p['departmentId'], p['departmentId']) if language == 'zh' else p['departmentId']} | {p['source']['sheet']}!{p['source']['cell']}]\n{p[field][language]}" for p in replies)
                assert assessment[target][language] == full_text, f"Lost complete current source paragraph: {code} {target} {language}"
        for entry in [master] + record["primary"] + record["supplementary"]:
            ws = live_workbooks[entry["source"]][entry["sheet"]]
            for coord, value in entry["cells"].items():
                assert ws[coord].value == value, f"Literal source-cell mismatch: {entry['source']} {entry['sheet']}!{coord}"
                literal_cells += 1
    assert chapter_counts == Counter(r["chapter"] for r in canonical) and len(chapter_counts) == 20
    assert chapter_counts["14"] == 22
    assert len([i for i in items.values() if i["frameworkId"] == "TMSA"]) == 288
    assert assessments["TMSA:2.4.2"]["response"] == "no" and assessments["TMSA:2.4.2"]["conformity"] == "partial"
    malformed = next(p for r in canonical for p in r["primary"] if p["raw_item"] == "aA.1.4")
    assert malformed["mapping_warning"] and prov["TMSA:canonical:1A.1.4"]["record"]["primary"] == next(r for r in canonical if r["item"] == "1A.1.4")["primary"]

    unit_index = {u["unit_id"]: u for u in units}
    assert len(unit_index) == len(units) == 570
    assert Counter(u["kind"] for u in units) == {"Expectation": 279, "Target": 291}
    referred_units = set()
    dry_rows = []
    evidence_paragraphs = 0
    gap_rows = []
    for area in mapping["areas"]:
        assert f"DRYBMS:{area['area']}" in chapters
        assert prov[f"DRYBMS:area:{area['area']}"]["record"] == area
        source_area = next(c for c in cells if c["area"] == area["area"])
        for row in area["requirements"]:
            dry_rows.append(row)
            actual = items[f"DRYBMS:{row['rid']}"]
            assert actual["analysisIndex"] and actual["code"] == row["rid"]
            assert actual["level"] == {"Basic": 1, "Intermediate": 2, "Advanced": 3, "Excellence": 4}[row["level"]]
            originals = [{"id": uid, "kind": unit_index[uid]["kind"], "en": unit_index[uid]["text_en"], "page": unit_index[uid]["page"]} for uid in row["source_unit_ids"]]
            assert actual["sourceOriginals"] == originals
            assert actual["kpi"] == {"zh": row["requirement_zh"], "en": "\n\n".join(u["en"] for u in originals)}
            suggested = []
            for c in source_area["cells"]:
                if c["page"] not in row["source_pages"]:
                    continue
                if area["area"] == 17 and c["page"] == 35 and c["level"] == "Basic" and row["level"] in {"Basic", "Intermediate"}:
                    raw_evidence = c["suggested_evidence_en"]
                    marker_at = raw_evidence.index("Verification of:")
                    suggested.append(raw_evidence[:marker_at] if row["level"] == "Basic" else raw_evidence[marker_at:])
                elif c["level"] == row["level"]:
                    suggested.append(c["suggested_evidence_en"])
            assert suggested, f"No same-page evidence: {row['rid']}"
            assert actual["guidance"]["en"] == "\n\n".join(suggested)
            assert actual["guidance"]["zh"] == "建議證據原文為英文；參考核對事項：" + row["attention_zh"]
            assert actual["crosswalk"] == {"tmsaItemCodes": row["tmsa_refs"], "match": row["match"], "difference": row["difference_zh"], "attention": row["attention_zh"]}
            assert all(f"TMSA:{x}" in items for x in row["tmsa_refs"])
            assert prov[f"DRYBMS:crosswalk:{row['rid']}"]["record"] == row
            assert prov[f"DRYBMS:crosswalk:{row['rid']}"]["sourceOriginals"] == originals
            assert actual.get("sourceGap", "") == row["gap_zh"]
            if row["gap_zh"]:
                assert actual["sourceGapResolved"] is False
                assert prov[f"DRYBMS:crosswalk:{row['rid']}"]["gapIds"]
                gap_rows.append(row["rid"])
            a = assessments[actual["id"]]
            assert a["response"] is None and not a["reviewed"]
            assert a["selfAssessment"] == a["evidence"] == {"en": "", "zh": ""}
            referred_units.update(row["source_unit_ids"])
            evidence_paragraphs += len(suggested)
    assert len(dry_rows) == len({r["rid"] for r in dry_rows}) == 412
    assert len(gap_rows) == 19 and len(mapping["gaps"]) == 20
    excluded = {x["unit_id"] for x in mapping["excluded_units"]}
    assert len(excluded) == 4 and not (excluded & referred_units)
    assert referred_units | excluded == set(unit_index), "Unaccounted source unit"
    assert prov["DRYBMS:gaps"]["record"] == mapping["gaps"]
    assert prov["DRYBMS:excluded-units"]["record"] == mapping["excluded_units"]
    assert prov["DRYBMS:original-units"]["record"] == units
    assert prov["DRYBMS:original-cells"]["record"] == cells
    dry_source = next(s for s in seed["sourceManifest"] if s["sourceKey"] == "drybms_v12")
    pdf_column_checks = verify_pdf_columns(private / dry_source["snapshotPath"], cells, units)
    assert all(s["levels"] == [None, None, None, None] for s in seed["dryScores"])

    history = load(tmsa / "out/historical_extra_fields.json")
    assert prov["TMSA:historical-extra-fields"]["record"] == history
    annotations = load(tmsa / "out/annotation_translations.json")
    assert prov["TMSA:annotations"]["record"] == annotations
    for pid, relative in [("TMSA:parsed-source-rows", "structured/source_rows.json"), ("TMSA:master-extra-cells", "structured/master_extras.json"),
                          ("TMSA:source-coverage", "structured/coverage.json"), ("TMSA:source-anomalies", "structured/anomalies.json"),
                          ("TMSA:template-exclusions", "structured/templates_excluded.json"), ("TMSA:registered-sources", "registered_sources.json")]:
        assert prov[pid]["record"] == load(tmsa / relative), f"Omitted auxiliary source archive: {pid}"
    cycles = load(tmsa / "out/cycle_review.json")
    assert seed["cycles"] == cycles["rows"]
    assert prov["TMSA:cycle-review-metadata"]["record"] == {"notes": cycles["notes"], "coverage": cycles["coverage"]}
    assert len(seed["cycles"]) == 146 and sum(bool(r["is_original"]) for r in seed["cycles"]) == 51
    quote_checks = 0
    for row in seed["cycles"]:
        assert f"TMSA:{row['item']}" in items
        for ref in row["basis_refs"]:
            assert live_workbooks[ref["source"]][ref["sheet"]][ref["field"]].value == ref["quote"], f"Cycle literal quote mismatch: {ref}"
            quote_checks += 1

    reports = {"mandeep": [load(tmsa / "out/audit_mandeep_early.json"), load(tmsa / "out/audit_mandeep_late.json")], "exxon": [load(tmsa / "out/audit_exxon.json")]}
    inspection_map = {i["id"]: i for i in seed["inspections"]}
    improvement_keys = [i["sourceKey"] for i in seed["improvements"]]
    assert len(improvement_keys) == len(set(improvement_keys))
    assert len(seed["improvements"]) == len({i["id"] for i in seed["improvements"]})
    assert len(prov["TMSA:external-finding-links"]["record"]) >= 105
    links = prov["TMSA:external-finding-links"]["record"]
    expected_improvements = set()
    inspection_counts = {}
    finding_quote_checks = 0
    for source_key, source_reports in reports.items():
        inspection = inspection_map[f"inspection:{source_key}:2025"]
        assert inspection["imported"] and inspection["status"] == "final"
        snap = inspection["snapshot"]
        assert snap["framework"] == frameworks["TMSA"]
        assert snap["chapters"] == [c for c in seed["chapters"] if c["frameworkId"] == "TMSA"]
        assert snap["items"] == [i for i in seed["items"] if i["frameworkId"] == "TMSA"]
        assert len(inspection["results"]) == 288
        assert {r["itemId"] for r in inspection["results"]} == {i["id"] for i in snap["items"]}
        assert all(r["response"] is None and not r["reviewed"] for r in inspection["results"])
        assert inspection["scoreNote"]["en"] and inspection["scoreNote"]["zh"]
        source_pdf = next(s for s in seed["sourceManifest"] if s["sourceKey"] == source_key)
        doc = pymupdf.open(private / source_pdf["snapshotPath"])
        findings = {f["id"]: f for f in inspection["findings"]}
        assert len(findings) == len(inspection["findings"])
        for report in source_reports:
            assert any(p["kind"] == "inspection-source-report" and p["record"] == report for p in seed["provenance"]), "Dropped audit coverage/general/no-finding metadata"
            for kind, entries in [("section", report["sections"]), ("general", report["general"])]:
                for entry in entries:
                    link = next(x for x in links if x["sourceKey"] == source_key and x["recordKind"] == kind and x["record"] == entry)
                    f = findings[link["findingId"]]
                    if kind == "section":
                        assert f["observation"] == {"en": entry["observation_en"], "zh": entry["observation_zh"]}
                        assert f["recommendation"] == {"en": entry["recommendation_en"], "zh": entry["recommendation_zh"]}
                        assert f["classification"] == entry["status"]
                        assert f["itemIds"] == [f"TMSA:{entry['item']}"]
                        assert entry["reported_heading"] == f["sources"][0]["rawItem"]
                        warning = entry.get("mapping_warning", "")
                        warning_zh = entry.get("mapping_warning_zh", "")
                        if warning:
                            assert warning in f["warning"] and warning_zh and warning_zh in f["warning"]
                        if entry.get("evidence_quote"):
                            page_text = "\n".join(doc[p - 1].get_text("text", sort=True) for p in entry["pdf_pages"])
                            assert compact(entry["evidence_quote"]) in compact(page_text), f"Nonliteral finding quote: {source_key} {entry['item']}"
                            finding_quote_checks += 1
                        should_improve = entry["status"] == "finding"
                    else:
                        assert f["observation"] == {"en": entry["text_en"], "zh": entry["text_zh"]}
                        assert f["itemIds"] == [], "Unnumbered general observations must not acquire guessed Items"
                        should_improve = entry["category"] in {"improvement_area", "element_summary"}
                    if should_improve:
                        expected_improvements.add(link["dedupeKey"])
                        improvement = next(i for i in seed["improvements"] if i["sourceKey"] == link["dedupeKey"])
                        assert improvement["inspectionId"] == inspection["id"] and improvement["findingId"] == f["id"]
                        assert improvement["source"] == "external" and improvement["status"] == "open" and improvement["percent"] == 0
                        assert improvement["closedAt"] is None and improvement["closedBy"] is None
                        assert improvement["assigneeIds"] == [] and improvement["targetDate"] == ""
                        assert improvement["itemIds"] == f["itemIds"]
                        for lang in ["en", "zh"]:
                            assert f["observation"][lang] in improvement["content"][lang]
                            assert f["recommendation"][lang] in improvement["content"][lang]
                        assert "Closure unverified" in improvement["sources"][0]["note"]
                    else:
                        assert link["dedupeKey"] not in improvement_keys, "Positive/guidance/no_finding promoted to deficiency"
        doc.close()
        inspection_counts[source_key] = {"findings": len(findings), "classifications": dict(Counter(f["classification"] for f in findings.values())), "unassessedResults": len(inspection["results"])}
    assert set(improvement_keys) == expected_improvements

    snapshot_manifest = load(private / "snapshot-manifest.json")
    assert prov["source-import:snapshot-manifest"]["record"] == snapshot_manifest
    assert snapshot_manifest["importedAt"] == stamp
    for snapshot in snapshot_manifest["files"]:
        target = private / snapshot["snapshotPath"]
        assert target.is_file() and sha(target) == snapshot["sha256"] and target.stat().st_size == snapshot["bytes"]
        if not snapshots_only:
            assert sha(Path(snapshot["sourcePath"])) == snapshot["sha256"], f"Changed extraction snapshot: {snapshot['sourcePath']}"
    assert len(snapshot_manifest["files"]) == len({x["snapshotPath"] for x in snapshot_manifest["files"]})
    expected_copy_paths = {x["snapshotPath"] for x in snapshot_manifest["files"]}
    actual_copy_paths = {p.relative_to(private).as_posix() for p in (private / "snapshots").rglob("*") if p.is_file()}
    assert actual_copy_paths == expected_copy_paths

    template_path = template_path or repo / "public/templates/tmsa-template.xlsx"
    assert template_path.is_file()
    template = load_workbook(template_path, data_only=False, keep_links=True)
    master = live_workbooks["master"]
    style = load(private / "tmsa-style-manifest.json")
    assert template.sheetnames == master.sheetnames
    assert len(template.sheetnames) == 20
    allowed = style["allowedCells"]
    actual_values = {}
    style_cells = 0
    for original_sheet, sanitized in zip(master, template):
        assert sanitized.sheet_state == original_sheet.sheet_state
        assert set(map(str, sanitized.merged_cells.ranges)) == set(map(str, original_sheet.merged_cells.ranges))
        assert str(sanitized.print_area) == str(original_sheet.print_area)
        assert sanitized.print_title_rows == original_sheet.print_title_rows
        assert sanitized.print_title_cols == original_sheet.print_title_cols
        assert sanitized.page_margins == original_sheet.page_margins
        assert sanitized.page_setup == original_sheet.page_setup
        assert sanitized.print_options == original_sheet.print_options
        assert set(sanitized.column_dimensions) == set(original_sheet.column_dimensions)
        for key, dim in original_sheet.column_dimensions.items():
            got = sanitized.column_dimensions[key]
            assert (got.width, got.min, got.max, got.hidden, got.outlineLevel) == (dim.width, dim.min, dim.max, dim.hidden, dim.outlineLevel)
        for key, dim in original_sheet.row_dimensions.items():
            got = sanitized.row_dimensions[key]
            assert (got.height, got.hidden, got.outlineLevel) == (dim.height, dim.hidden, dim.outlineLevel)
        values = {}
        for coord, cell in sanitized._cells.items():
            if isinstance(cell, MergedCell):
                continue
            assert not cell.comment and not cell.hyperlink
            assert cell.data_type != "f"
            if cell.value is not None:
                values[cell.coordinate] = cell.value
        assert values == allowed[sanitized.title], f"Non-whitelisted template content: {sanitized.title}"
        actual_values[sanitized.title] = values
        for coord, cell in original_sheet._cells.items():
            if isinstance(cell, MergedCell):
                continue
            got = sanitized.cell(*coord)
            assert all(copy.copy(getattr(got, prop)) == copy.copy(getattr(cell, prop)) for prop in ["font", "fill", "border", "alignment", "protection"]), f"Changed style: {sanitized.title}!{cell.coordinate}"
            assert got.number_format == cell.number_format
            style_cells += 1
        assert not sanitized._images and not sanitized._charts and not sanitized.tables
        assert len(sanitized.data_validations.dataValidation) == 0
        assert all(getattr(getattr(sanitized, name), part).text is None for name in ["oddHeader", "oddFooter", "evenHeader", "evenFooter", "firstHeader", "firstFooter"] for part in ["left", "center", "right"])
    assert not template._external_links and not template.defined_names
    template_ids = [value for sheet in actual_values.values() for value in sheet.values() if isinstance(value, str) and re.fullmatch(r"\d+[A-Z]?\.[1-4]\.\d+", value)]
    assert len(template_ids) == len(set(template_ids)) == 288
    assert set(template_ids) == {r["item"] for r in canonical}
    excluded_texts = {str(c.value) for ws in master for c in ws._cells.values() if not isinstance(c, MergedCell) and c.value is not None and c.coordinate not in allowed[ws.title] and len(str(c.value)) >= 12}
    with zipfile.ZipFile(template_path) as z:
        assert z.testzip() is None
        names = z.namelist()
        assert not any(any(word in name.lower() for word in ["comment", "externallink", "vml", "connections", "customxml", "drawing", "embedding", "printersettings"]) for name in names)
        # This ZIP was generated locally by openpyxl from the hash-bound source,
        # not uploaded arbitrary XML. Reject declarations before stdlib parsing;
        # ElementTree does not resolve external entities in these OOXML parts.
        for name in names:
            if name.endswith(".xml"):
                assert b"<!DOCTYPE" not in z.read(name) and b"<!ENTITY" not in z.read(name)
        joined_text = "\n".join("".join(ET.fromstring(z.read(n)).itertext()) for n in names if n.endswith(".xml"))
        assert not any(text in joined_text for text in excluded_texts), "Original narrative leaked in public ZIP"
        assert not any(b'TargetMode="External"' in z.read(n) for n in names if n.endswith(".rels"))
    assert template.properties.creator in (None, "") and template.properties.lastModifiedBy in (None, "")
    assert len(template.custom_doc_props) == 0
    # Durable private data must remain ignored, and not be a public asset.
    ignored = subprocess.run(["git", "-C", str(repo), "check-ignore", "private-seed/platform-seed.json"], capture_output=True, text=True)
    assert ignored.returncode == 0, "private-seed is not gitignored"
    public_leaks = []
    for root in [repo / "public", repo / "dist"]:
        if not root.exists():
            continue
        for asset in root.rglob("*"):
            if not asset.is_file() or asset == template_path or asset.name == "tmsa-template.xlsx":
                continue
            blob = asset.read_bytes()
            if b'"source-import"' in blob and b'"sourceManifest"' in blob:
                public_leaks.append(str(asset))
            # Exact substantive English paragraphs, not generic labels/logo.
            if any(text.encode("utf-8") in blob for text in excluded_texts if len(text) > 80):
                public_leaks.append(str(asset))
    assert not public_leaks, f"Private source text in public/dist: {public_leaks}"
    template_hash_before_native = sha(template_path)
    native = native_excel(template_path) if run_native else {"verified": False, "note": "Not requested; use --native-excel on Windows with Excel installed."}
    assert sha(template_path) == template_hash_before_native, "Read-only native verification modified the template"
    if not snapshots_only:
        assert all(sha(Path(s["path"])) == s["sha256"] for s in seed["sourceManifest"]), "An original changed during verification"
    for workbook in live_workbooks.values():
        workbook.close()
    template.close()
    return {
        "passed": True, "verifiedAt": datetime.now(timezone.utc).isoformat(), "importedAt": stamp,
        "scope": "durable snapshots" if snapshots_only else "live originals + durable snapshots",
        "counts": {"frameworks": 2, "tmsaChapters": 20, "drybmsChapters": 30, "tmsaItems": 288, "drybmsAnalysisRows": 412,
                   "assessments": 700, "registryTexts": len(registry), "additionalRegistryTexts": len(extra_registry),
                   "registryUses": sum(len(x["uses"]) for x in all_registry.values()), "primaryResponses": sum(len(r["primary"]) for r in canonical),
                   "supplementaryResponses": sum(len(r["supplementary"]) for r in canonical), "copiedTemplateReferences": sum(len(r["template_sources"]) for r in canonical),
                   "historicalExtraFields": len(history), "pdfAnnotationRecords": len(annotations["pdf_annotations"]), "distinctExcelComments": len(annotations["excel_comments"]),
                   "cycles": len(seed["cycles"]), "originalCycleRows": 51, "cycleLiteralQuotes": quote_checks,
                   "sourceUnits": 570, "expectationUnits": 279, "targetUnits": 291, "referredUnits": len(referred_units), "excludedUnits": 4,
                   "gapRecords": 20, "affectedRows": 19, "inspections": 2, "improvements": len(seed["improvements"]),
                   "snapshotFiles": len(snapshot_manifest["files"]), "literalExcelCellsChecked": literal_cells, "auditEvidenceQuotes": finding_quote_checks,
                   "literalDrybmsPdfColumns": pdf_column_checks},
        "tmsaChapterCounts": dict(chapter_counts), "conflictingSourceResponses": conflicts, "inspectionCounts": inspection_counts,
        "originalSourceHashes": original_hashes, "template": {"path": str(template_path), "sha256": sha(template_path), "worksheets": 20, "literalItemCodes": 288,
        "nonemptyWhitelistedCells": sum(len(x) for x in actual_values.values()), "preservedStyleCells": style_cells,
        "comments": 0, "externalLinks": 0, "formulas": 0, "originalNarrativeLeaks": 0, "nativeExcel": native},
        "publicPrivateDataLeaks": [], "errors": [],
        "unresolved": ["Source gaps: 20 records affect 19 DRYBMS analysis rows; no text guessed.",
                       "TMSA 2.4.2 retains crew No / management Yes as No/partial.",
                       "All imported current assessments require fresh review; historical inspections lack complete per-item scores.",
                       "General audit themes remain unlocated (no guessed Item IDs); all external improvement closure is unverified."]
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--repo", type=Path, default=DEFAULT_REPO)
    parser.add_argument("--snapshots-only", action="store_true", help="For relocation: verify private copies without requiring the original absolute paths")
    parser.add_argument("--native-excel", action="store_true", help="Open the public template read-only in an owned Microsoft Excel COM instance")
    parser.add_argument("--no-write", action="store_true", help="Read-only validation, useful for negative tests")
    parser.add_argument("--seed", type=Path, help="Alternate test seed; requires --no-write")
    parser.add_argument("--template", type=Path, help="Alternate test template; requires --no-write")
    args = parser.parse_args()
    if (args.seed or args.template) and not args.no_write:
        parser.error("Alternate test inputs require --no-write; never replace the official verification report with a fixture.")
    try:
        report = verify(args.repo.resolve(), args.snapshots_only, args.native_excel, args.seed, args.template)
    except (AssertionError, OSError, KeyError, ValueError, StopIteration, subprocess.SubprocessError) as error:
        import traceback
        frame = traceback.extract_tb(error.__traceback__)[-1]
        detail = f"{type(error).__name__}: {str(error) or frame.line} ({Path(frame.filename).name}:{frame.lineno})"
        report = {"passed": False, "verifiedAt": datetime.now(timezone.utc).isoformat(), "errors": [detail]}
    if not args.no_write:
        output = args.repo / "private-seed/source-verification.json"
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["passed"] else 1


if __name__ == "__main__":
    sys.exit(main())
