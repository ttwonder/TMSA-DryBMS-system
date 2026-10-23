"""Owned read-only Excel: real schema, AutoFit and filter/sort associations."""
import sys,json,hashlib,shutil,uuid,gc,subprocess
from pathlib import Path
from collections import Counter

def digest(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def worker(entry):
 import win32com.client,win32process,win32api,win32con,win32event
 path=Path(entry['path']).resolve();before=digest(path);copy=path.with_name(path.stem+'-native-'+uuid.uuid4().hex[:8]+'.xlsx');shutil.copy2(path,copy)
 existing=set(win32process.EnumProcesses());app=win32com.client.DispatchEx('Excel.Application');app.Visible=False;app.DisplayAlerts=False;app.EnableEvents=False;app.AutomationSecurity=3
 _,pid=win32process.GetWindowThreadProcessId(app.Hwnd);assert pid not in existing
 handle=win32api.OpenProcess(win32con.SYNCHRONIZE|win32con.PROCESS_TERMINATE,False,pid);book=sheet=rng=sort=None
 result={'source':str(path),'sourceSha256':before,'ownedPid':pid,'copy':str(copy),'sheets':[]}
 try:
  book=app.Workbooks.Open(str(copy),UpdateLinks=0,ReadOnly=True,CorruptLoad=0);assert book.ReadOnly
  for n in range(1,book.Worksheets.Count+1):
   sheet=book.Worksheets(n);rng=sheet.AutoFilter.Range;first=int(rng.Row)+1;last=int(rng.Row+rng.Rows.Count-1);cols=int(rng.Columns.Count)
   sr={'name':sheet.Name,'columns':cols,'first':first,'last':last};result['sheets'].append(sr)
   if last<first:continue
   vals=lambda:[tuple(sheet.Range(sheet.Cells(i,1),sheet.Cells(i,cols)).Value2[0]) for i in range(first,last+1)]
   original=vals();heights=[sheet.Rows(i).RowHeight for i in range(first,last+1)]
   sheet.Range(sheet.Cells(first,1),sheet.Cells(last,cols)).Rows.AutoFit();needed=[sheet.Rows(i).RowHeight for i in range(first,last+1)]
   sr.update(declared=heights,autofit=needed,clipped=[first+i for i,(a,b) in enumerate(zip(heights,needed)) if a and b>a+1],excess=[round(a-b,2) if a else 0 for a,b in zip(heights,needed)])
   sr['excessWholeLines']=[first+i for i,x in enumerate(sr['excess']) if x>=15.4]
   sr['columnWidths']=[sheet.Columns(i).Width for i in range(1,cols+1)]
   sr['footer']={'left':sheet.PageSetup.LeftFooter,'center':sheet.PageSetup.CenterFooter}
   assert sr['footer']['left'].startswith('&12') and sr['footer']['center'].startswith('&12'),sr['footer']
   # Render original row heights, not an AutoFit-repaired workbook.
   for i,h in enumerate(heights,first):
    if h is not None:sheet.Rows(i).RowHeight=h
   pdf=copy.with_name(copy.stem+'-'+str(n)+'.pdf');sheet.ExportAsFixedFormat(0,str(pdf));sr['render']=str(pdf)
   key=2 if entry['request'].get('filters',{}).get('reportVariant')=='remarks' else 1
   nonempty=[row for row in original if row[key-1] is not None]
   if nonempty:
    value=nonempty[0][key-1];rng.AutoFilter(Field=key,Criteria1=value)
    visible=[row for i,row in enumerate(vals(),first) if not sheet.Rows(i).Hidden]
    assert Counter(visible)==Counter(row for row in original if row[key-1]==value);sheet.ShowAllData();sr['filter']='PASS'
    header=tuple(sheet.Range(sheet.Cells(first-1,1),sheet.Cells(first-1,cols)).Value2[0])
    for order in [2,1]:
     sort=sheet.Sort;sort.SortFields.Clear();sort.SortFields.Add(Key=sheet.Range(sheet.Cells(first,key),sheet.Cells(last,key)),SortOn=0,Order=order,DataOption=0);sort.SetRange(rng);sort.Header=1;sort.MatchCase=False;sort.Orientation=1;sort.Apply()
     assert Counter(vals())==Counter(original);assert tuple(sheet.Range(sheet.Cells(first-1,1),sheet.Cells(first-1,cols)).Value2[0])==header
     sheet.Range(sheet.Cells(first,1),sheet.Cells(last,cols)).Rows.AutoFit()
    sr['sortAssociationsAndExplicitAutoFit']='PASS'
  result['nativeAcceptance']='PASS';result['initialHeights']='PASS' if not any(s.get('clipped') or s.get('excessWholeLines') for s in result['sheets']) else 'FAIL'
 finally:
  sort=rng=sheet=None
  if book is not None:book.Close(SaveChanges=False)
  book=None;app.Quit();app=None;gc.collect();stopped=win32event.WaitForSingleObject(handle,8000)==win32event.WAIT_OBJECT_0
  if not stopped:win32api.TerminateProcess(handle,0);result['forcedOwnedStop']=True;stopped=win32event.WaitForSingleObject(handle,5000)==win32event.WAIT_OBJECT_0
  handle.Close();result['ownedExcelStopped']=stopped;assert stopped;assert digest(path)==before and digest(copy)==before;result['sourceAndCopyUnchanged']=True
 return result

def main():
 if '--worker' in sys.argv:print(json.dumps(worker(json.loads(sys.stdin.read())),ensure_ascii=True));return
 receipt=Path(sys.argv[1]);data=json.loads(receipt.read_text(encoding='utf8'));out=receipt.parent/('native-refinement-'+uuid.uuid4().hex[:8]+'.json');result={'inputReceipt':str(receipt),'sha256':digest(receipt),'artifacts':[],'status':'RUNNING'}
 try:
  for e in data['downloads']:
   if not e['path'].endswith('.xlsx'):continue
   from openpyxl import load_workbook
   book=load_workbook(e['path'],rich_text=True)
   bad=[]
   for sh in book:
    for row in sh.iter_rows(min_row=2):
     for cell in row:
      if cell.value is None:continue
      if cell.font.sz!=12:bad.append((sh.title,cell.coordinate,cell.font.sz))
      if isinstance(cell.value,list):
       for run in cell.value:
        if hasattr(run,'font') and run.font.sz!=12:bad.append((sh.title,cell.coordinate,'run',run.font.sz))
   assert not bad,bad
   r=subprocess.run([sys.executable,__file__,'--worker'],input=json.dumps(e),capture_output=True,text=True,encoding='utf8',timeout=150)
   if r.returncode:result['artifacts'].append({'path':e['path'],'status':'FAIL','error':r.stderr});raise RuntimeError(r.stderr)
   result['artifacts'].append(json.loads(r.stdout));out.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
  result['status']='PASS' if all(a['initialHeights']=='PASS' for a in result['artifacts']) else 'FAIL'
 except Exception as ex:result['status']='FAIL';result['failure']=str(ex)
 finally:out.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8');print(json.dumps({'status':result['status'],'receipt':str(out),'artifacts':len(result['artifacts'])}));sys.exit(0 if result['status']=='PASS' else 1)
if __name__=='__main__':main()
