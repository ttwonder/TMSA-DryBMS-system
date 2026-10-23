"""Verify real report-selection downloads; native Excel uses owned read-only copies."""
import gc, hashlib, json, re, shutil, subprocess, sys, uuid
from collections import Counter
from pathlib import Path
import fitz
import openpyxl

def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def flat(value):
    return re.sub(r'\s+', '', str(value or ''))

def native(entry):
    import win32api, win32con, win32event, win32process, win32com.client
    path=Path(entry['path']); before=sha(path)
    copy=path.with_name(path.stem+'-owned-'+uuid.uuid4().hex[:8]+'.xlsx');shutil.copy2(path,copy)
    existing=set(win32process.EnumProcesses());app=win32com.client.DispatchEx('Excel.Application')
    app.Visible=False;app.DisplayAlerts=False;app.EnableEvents=False;app.AutomationSecurity=3
    _,pid=win32process.GetWindowThreadProcessId(app.Hwnd);assert pid not in existing
    handle=win32api.OpenProcess(win32con.SYNCHRONIZE|win32con.PROCESS_TERMINATE,False,pid)
    result={'ownedPid':pid,'sourceSha256':before,'sheets':[]};book=sheet=rng=sort=None
    try:
        book=app.Workbooks.Open(str(copy.resolve()),UpdateLinks=0,ReadOnly=True,CorruptLoad=0)
        assert book.ReadOnly
        for n in range(1,book.Worksheets.Count+1):
            sheet=book.Worksheets(n);rows=int(sheet.UsedRange.Rows.Count);cols=int(sheet.UsedRange.Columns.Count)
            part={'name':sheet.Name,'rows':rows,'columns':cols}
            operational=entry['kind']=='improvements' or entry['request'].get('filters',{}).get('reportVariant')=='remarks'
            if operational:
                assert sheet.AutoFilterMode
                rng=sheet.AutoFilter.Range
                assert int(rng.Row)==4 and int(rng.Rows.Count)==max(1,rows-3)
                saved=sheet.Rows(4).RowHeight;sheet.Rows(4).AutoFit();needed=sheet.Rows(4).RowHeight
                assert needed<=saved+1,('header clipping',sheet.Name,saved,needed)
                sheet.Rows(4).RowHeight=saved;part['headerHeight']={'saved':saved,'nativeRequired':needed}
                if rows>4:
                    original=[tuple(sheet.Range(sheet.Cells(i,1),sheet.Cells(i,cols)).Value2[0]) for i in range(5,rows+1)]
                    key=1 if entry['kind']=='improvements' else 2
                    rng.AutoFilter(Field=key,Criteria1=original[0][key-1]);visible=[tuple(sheet.Range(sheet.Cells(i,1),sheet.Cells(i,cols)).Value2[0]) for i in range(5,rows+1) if not sheet.Rows(i).Hidden]
                    assert Counter(visible)==Counter(v for v in original if v[key-1]==original[0][key-1]);sheet.ShowAllData()
                    header=tuple(sheet.Range(sheet.Cells(4,1),sheet.Cells(4,cols)).Value2[0])
                    for order in (2,1):
                        sort=sheet.Sort;sort.SortFields.Clear();sort.SortFields.Add(Key=sheet.Range(sheet.Cells(5,key),sheet.Cells(rows,key)),SortOn=0,Order=order,DataOption=0);sort.SetRange(rng);sort.Header=1;assert sort.Header==1;sort.Orientation=1;sort.Apply()
                        after=[tuple(sheet.Range(sheet.Cells(i,1),sheet.Cells(i,cols)).Value2[0]) for i in range(5,rows+1)]
                        assert Counter(after)==Counter(original);assert [v[key-1] for v in after]==sorted([v[key-1] for v in original],reverse=order==2)
                        assert tuple(sheet.Range(sheet.Cells(4,1),sheet.Cells(4,cols)).Value2[0])==header
                    heights=[sheet.Rows(i).RowHeight for i in range(5,rows+1)];sheet.Range(sheet.Cells(5,1),sheet.Cells(rows,cols)).Rows.AutoFit();needed=[sheet.Rows(i).RowHeight for i in range(5,rows+1)]
                    assert all(a<=b+1 for a,b in zip(needed,heights)),('body clipping',sheet.Name)
                    for i,h in enumerate(heights,5):sheet.Rows(i).RowHeight=h
                    part.update(filterAndBothSortDirections='PASS',bodyRequiredMax=max(needed),savedHeightMax=max(heights))
            # A disposable native print of the new main-list entry: actual saved widths/fonts.
            if entry['name']=='main-single-xlsx' and n==1:
                sheet.PageSetup.PrintArea='A1:H6';sheet.PageSetup.Zoom=False;sheet.PageSetup.FitToPagesWide=1;sheet.PageSetup.FitToPagesTall=False
                pdf=path.with_name('native-main-metadata.pdf');sheet.ExportAsFixedFormat(0,str(pdf.resolve()));part['nativeRender']=str(pdf)
            result['sheets'].append(part)
        result['CorruptLoad0ReadOnly']='PASS'
    finally:
        sort=rng=sheet=None
        if book is not None:book.Close(SaveChanges=False)
        book=None;app.Quit();app=None;gc.collect()
        stopped=win32event.WaitForSingleObject(handle,8000)==win32event.WAIT_OBJECT_0
        if not stopped:win32api.TerminateProcess(handle,0);result['forcedOwnedStop']=True;stopped=win32event.WaitForSingleObject(handle,5000)==win32event.WAIT_OBJECT_0
        win32api.CloseHandle(handle);result['ownedProcessStopped']=stopped
        assert stopped and sha(path)==before
    return result

def verify(receipt_path):
    path=Path(receipt_path);receipt=json.loads(path.read_text(encoding='utf-8'));assert receipt['status']=='PASS'
    result={'status':'RUNNING','layer':'Saved browser downloads + independent parser + native Excel COM','sourceReceiptSha256':sha(path),'verifierSha256':sha(__file__),'artifacts':[]}
    target=path.with_name('artifacts.json')
    try:
        for entry in receipt['downloads']:
            assert sha(entry['path'])==entry['sha256'] and sha(entry['canonical'])==entry['canonicalSha256']
            state=json.loads(Path(entry['canonical']).read_text(encoding='utf-8'));part={'name':entry['name'],'sha256':entry['sha256']}
            if entry['format']=='pdf':
                doc=fitz.open(entry['path']);text='\n'.join(p.get_text() for p in doc);compact=flat(text)
                for page in doc:
                    spans=[s for b in page.get_text('dict')['blocks'] if 'lines' in b for line in b['lines'] for s in line['spans']]
                    assert spans
                    for s in spans:assert s['size']>=13.99 and s['bbox'][0]>=32 and s['bbox'][2]<=page.rect.width-32
                if entry['kind']=='framework':
                    remark=entry['request'].get('filters',{}).get('reportVariant')=='remarks';prefixes=['REMARK_MARKER_'] if remark else ['KPI_MARKER_','GUIDANCE_MARKER_','ASSESSMENT_MARKER_','EVIDENCE_MARKER_']
                    for item in state['items']:
                        for prefix in prefixes:assert (flat(prefix+item['id']) in compact)==(item['id'] in entry['expectedItems']),(entry['name'],prefix,item['id'])
                    if not remark:
                        filtered=bool(entry['request'].get('filters',{}).get('departmentIds'))
                        if filtered:assert '不按部門重算' in compact and '部門自評摘錄' in compact
                        else:assert '完整自評報告' in compact and '不按部門重算' not in compact
                if entry['kind']=='improvements':
                    for marker in ['COMPLETE_ACTION_0','PLAN_MARKER_0','PROGRESS_MARKER_COMPLETE','MILESTONE_MARKER_COMPLETE']:assert marker in compact
                    assert 'OTHERAction' not in compact
                image=Path(entry['path']).with_suffix('.png');doc[0].get_pixmap(matrix=fitz.Matrix(1.3,1.3)).save(image);part.update(pages=len(doc),image=str(image),textAndFontFloor='PASS');doc.close()
            else:
                book=openpyxl.load_workbook(entry['path']);part['sheets']=book.sheetnames
                if entry['kind']=='improvements':
                    sheet=book.worksheets[0];assert sheet.freeze_panes=='A5' and sheet.auto_filter.ref==f'A4:S{sheet.max_row}'
                    assert not any(m.min_row>=4 for m in sheet.merged_cells.ranges)
                    assert sheet.column_dimensions['N'].width>sheet.column_dimensions['A'].width
                    rows=list(sheet.iter_rows(min_row=5,values_only=True))
                    assert set(r[0] for r in rows)==set(range(1,len(entry['request']['ids'])+1))
                    for number,id in enumerate(entry['request']['ids'],1):
                        original=next(r for r in state['improvements'] if r['id']==id);selected=sorted([r for r in rows if r[0]==number],key=lambda r:r[15]);content=''.join(r[13] or '' for r in selected);plan=''.join(r[14] or '' for r in selected)
                        assert flat(original['content']['en']) in flat(content) and flat(original['content']['zh']) in flat(content);assert flat(plan)==flat(original['plan'])
                        for r in selected:assert r[11].strftime('%Y-%m-%d')==original['targetDate']
                worker=path.with_name(entry['name']+'-native-request.json');worker.write_text(json.dumps(entry),encoding='utf-8')
                process=subprocess.run([sys.executable,__file__,'--native',str(worker)],capture_output=True,text=True,encoding='utf-8',timeout=120)
                assert process.returncode==0,process.stderr or process.stdout;part['native']=json.loads(process.stdout)
                for sheet in part['native']['sheets']:
                    if sheet.get('nativeRender'):
                        d=fitz.open(sheet['nativeRender']);image=Path(sheet['nativeRender']).with_suffix('.png');d[0].get_pixmap(matrix=fitz.Matrix(2,2)).save(image);part['nativeImage']=str(image);d.close()
            result['artifacts'].append(part);target.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
        result['status']='PASS';result['artifactCount']=len(result['artifacts']);assert result['artifactCount']==len(receipt['downloads'])
    except Exception as error:result.update(status='FAIL',failure=str(error));raise
    finally:target.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'status':result['status'],'artifactCount':result['artifactCount'],'receipt':str(target)},ensure_ascii=False))

if __name__=='__main__':
    if sys.argv[1]=='--native':print(json.dumps(native(json.loads(Path(sys.argv[2]).read_text(encoding='utf-8'))),ensure_ascii=False))
    else:verify(sys.argv[1])
